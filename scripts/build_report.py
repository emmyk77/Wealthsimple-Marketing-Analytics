"""Create the brief's two-page recommendation report from reconciled metrics."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
R=json.loads((ROOT/'data/analysis/findings.json').read_text());t=R['totals'];q=R['quality']
d=Document();sec=d.sections[0];sec.page_height=Inches(11.7);sec.page_width=Inches(8.3)
sec.top_margin=sec.bottom_margin=Inches(.62);sec.left_margin=sec.right_margin=Inches(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 s=d.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(6);d.styles['Normal'].paragraph_format.line_spacing=1.06
d.styles['Title'].font.size=Pt(22);d.styles['Title'].paragraph_format.space_after=Pt(5)
d.styles['Subtitle'].font.size=Pt(11)
d.styles['Heading 1'].font.size=Pt(13);d.styles['Heading 1'].paragraph_format.space_before=Pt(9);d.styles['Heading 1'].paragraph_format.space_after=Pt(4)
for style in d.styles:
 for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
def p(x):d.add_paragraph(x)
def h(x):d.add_heading(x,level=1)
def table(rows,widths):
 tb=d.add_table(rows=0,cols=len(rows[0]));tb.autofit=False
 for col,width in zip(tb.columns,widths):col.width=Inches(width)
 for vals in rows:
  cells=tb.add_row().cells
  for j,v in enumerate(vals):
   cells[j].width=Inches(widths[j]);cells[j].text=str(v)
   for para in cells[j].paragraphs:
    para.paragraph_format.space_after=Pt(4);para.paragraph_format.space_before=Pt(4)
    for r in para.runs:r.font.size=Pt(10)
 for i,row in enumerate(tb.rows):
  for cell in row.cells:
   pr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'17324D' if i==0 else ('F2F5F7' if i%2==0 else 'FFFFFF'));pr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   pr.append(borders)
   if i==0:
    for r in cell.paragraphs[0].runs:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
 d.add_paragraph().paragraph_format.space_after=Pt(1)
 return tb
d.add_paragraph('Marketing efficiency and customer conversion',style='Title')
d.add_paragraph('Wealthsimple Canada training case | Recommendation report | 28 October 2026',style='Subtitle')
p('Prepared by Emmanuel F. Kebodi. Fictional training data; these are not actual Wealthsimple business results.')
h('Executive recommendation')
p('Prioritize measurement repair and controlled acquisition tests before expanding budgets. The comparable population records CAD 426,714 in marketing cost, 13,700 signups and 106 paying customers: a 0.77% signup-to-paying rate and CAD 4,026 CAC. Paid Search contributes 69 of those customers. Email and Referral have lower observed CAC, but their small customer counts do not establish that they can absorb large budgets efficiently.')
h('Scope and data quality')
p('All 2,000 original campaign records remain available. The headline analysis includes 1,671 records with usable cost and a complete, nonnegative, sequential funnel through KYC. It excludes 329 records from comparison, representing CAD 82,177 of known cost. The cleaned dataset contains CAD 508,892 of known cost overall; six records still lack usable cost values. These figures are partial recorded amounts, not a certified total expenditure.')
p('The workflow standardizes 47 channel variants into 11 channels and 31 segment variants into eight segments. It leaves invalid numbers null, reconstructs 72 total-cost values only from complete components, and preserves issue logs. There are 102 missing or invalid dates, 107 date versus campaign-month conflicts and 286 campaign IDs associated with multiple names. Detailed campaigns use ID plus name; campaign families use channel plus name.')
p('Three source records each show 100 paying customers with zero funded accounts. Such anomalies invalidate the earlier 607-customer headline. The strict funnel assumes the counts describe the same cohort, which requires source-owner confirmation. Sensitivity analysis gives 108 customers for consistent funnels and 96 when the headline population also requires an unconflicted date. No unique-customer identifier is available.')
h('Spend and channel performance')
table([['Channel','Cost CAD','Paying','CAC CAD'],['Paid Search','60,857','69','882'],['Referral','4,503','13','346'],['Organic Search','6,739','9','749'],['Email','2,939','9','327'],['Affiliate','14,596','6','2,433']],[2.6,1.45,1.1,1.35])
p('Six other channels record no paying customers in the eligible population. Across 19 campaign families, CAD 337,080 (79.0% of eligible cost) is associated with zero recorded paying customers. This is spend requiring investigation, not proven waste: awareness objectives, delayed conversion and attribution quality must be checked. Trade Account Search records 36 customers at CAD 473 CAC; it is a stronger candidate for a controlled test than an unconditional scaling decision.')
d.add_page_break()
h('Conversion and audience findings')
p('The comparable funnel contains 4,557,933 impressions, 136,235 clicks, 13,700 signups, 1,673 funded accounts and 106 paying customers. Step conversion is 2.99%, 10.06%, 12.21% and 6.34%, respectively. Impressions to clicks has the largest absolute and percentage drop-off. Among account holders, 12,027 signups do not reach funding and 1,567 funded accounts do not reach payment. These observations support testing onboarding and activation; they do not prove the cause of abandonment or that CTR is healthy.')
p('First-time Investors record 19 paying customers and CAD 2,685 CAC, compared with five and CAD 10,999 for Parents. Quebec records 17 paying customers at CAD 2,655 CAC; Ontario records six at CAD 6,869. Desktop records 30 paying customers and 1.09% signup-to-paying conversion, versus 15 and 0.59% on iOS. Mortgage has CAD 62,758 cost but only three paying customers; product sales cycles and campaign mix could explain some differences. These are descriptive comparisons with small counts, not causal rankings.')
h('Revenue and return interpretation')
p('Known first-month revenue is CAD 4,689 on 1,646 eligible records. On matched rows, revenue divided by ad spend is 0.0135x, while revenue divided by total marketing cost is 0.0112x. The matched revenue-minus-marketing-cost gap is CAD −414,869. This is not gross margin or lifetime profit: operating costs, margin, retention, attribution windows and later revenue are unavailable. The 1,886 zero-revenue source rows are distinct from 33 missing or invalid revenue values. Zero revenue alone does not establish a tracking failure.')
h('Recommended actions and measurement')
p('1. Repair data capture first. The analytics owner should reconcile the three 100-customer anomalies, negative values, campaign identity collisions and date conflicts. Require stable campaign identifiers, numeric validation and documented cohort windows. Monitor excluded-record share and missing revenue separately. Retain a raw-to-clean audit trail.')
p('2. Review zero-output campaign families with channel owners. Start with Investing Explained Pre-roll (CAD 25,150), Retargeting Carousel (CAD 24,491) and Finance News Display (CAD 23,093). For conversion objectives, consider a pause or creative reset only after checking attribution lag and data validity. Evaluate awareness activity against its stated objective rather than automatically calling it waste.')
p('3. Run a small, capped acquisition test. Compare Trade Account Search, Referral and Email using matched audiences, time windows and attribution rules. Track incremental paying customers, CAC and retention with a holdout where practical. Historical CAC is descriptive; neither a guaranteed CAC reduction nor additional revenue can be estimated reliably from these data.')
p('4. Test account activation. Product and lifecycle teams should test clearer funding instructions, KYC reminders and timely follow-up after funding. Measure signup-to-funded and funded-to-paying rates separately, with a defined observation window. Break down results by device, product and segment before reallocating regional budgets.')
p('5. Agree commercial targets with Finance. Set CAC limits from contribution margin, expected retention and payback requirements, rather than adopting the earlier arbitrary CAD 400 threshold. Review the seven dashboard areas monthly: data quality, spend and CAC, funnel, channel and campaign, revenue, audience and region, and budget optimization.')
p('Basis: supplied Power_BI_project.xlsx and the reproducible cleaned analysis in the project repository. Scenario dates span 2025–2026; future-dated records are not evidence of realized company performance. Monthly comparisons exclude CAD 39,959 of eligible cost with missing or conflicting dates. Validate Power BI refresh totals before submitting the final PBIX.')
d.core_properties.author='Emmanuel F. Kebodi';d.core_properties.title='Marketing efficiency and customer conversion'
out=ROOT/'reports/Wealthsimple_Recommendation_Report.docx';out.parent.mkdir(exist_ok=True);d.save(out);print(out)
